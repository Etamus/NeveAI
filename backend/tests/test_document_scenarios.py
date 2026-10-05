import asyncio
import io
import json
import random
import unittest
from unittest.mock import AsyncMock

from neveai.utils.document_edits import (
    inspect_document,
    apply_document_edits,
    DocumentEditError,
    TEXT_FORMATS,
)
from neveai.utils.document_edit_planner import (
    propose_document_edits,
    quoted_replacement,
)
from neveai.utils.document_text import decode_document_text
from neveai.utils.document_validation import validate_document_output
from neveai.utils.generated_files import (
    build_generated_file,
    MIME_TYPES,
    GeneratedFileError,
)


def saved(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class TextScenarios(unittest.TestCase):
    def test_all_native_text_formats_edit_and_validate_without_changing_surroundings(
        self,
    ):
        fixtures = {
            "json": '{"nome":"Marcos","valor":1}\n',
            "xml": "<doc><nome>Marcos</nome></doc>\n",
            "yaml": "nome: Marcos\nvalor: 1\n",
            "yml": "nome: Marcos\nvalor: 1\n",
            "py": 'def nome():\n    return "Marcos"\n',
            "csv": "Nome,Valor\nMarcos,1\n",
            "srt": "1\n00:00:00,000 --> 00:00:01,000\nMarcos\n",
        }
        for fmt in TEXT_FORMATS:
            with self.subTest(format=fmt):
                source = fixtures.get(fmt, "Antes\nMarcos\nDepois\n").encode()
                snapshot = inspect_document(source, fmt)
                changes = quoted_replacement(snapshot, 'Troque "Marcos" por "Mateus"')
                data, _ = apply_document_edits(snapshot, changes)
                self.assertEqual(data, source.replace(b"Marcos", b"Mateus"))
                self.assertEqual(snapshot.data, source)
                validate_document_output(data, fmt)

    def change(self, original, objective):
        snapshot = inspect_document(original, "txt")
        return apply_document_edits(snapshot, quoted_replacement(snapshot, objective))

    def test_utf16_utf32_and_boms_survive_edits(self):
        for encoding, bom in [
            ("utf-16-le", b"\xff\xfe"),
            ("utf-16-be", b"\xfe\xff"),
            ("utf-32-le", b"\xff\xfe\0\0"),
            ("utf-32-be", b"\0\0\xfe\xff"),
            ("utf-8", b"\xef\xbb\xbf"),
        ]:
            with self.subTest(encoding=encoding):
                source = bom + "Nome: Marcos\r\nM\u00fasica mantida\r\n".encode(
                    encoding
                )
                data, _ = self.change(source, 'Troque "Marcos" por "Mateus"')
                self.assertEqual(
                    data,
                    bom + "Nome: Mateus\r\nM\u00fasica mantida\r\n".encode(encoding),
                )
                json.dumps(inspect_document(data, "txt").inventory)

    def test_cp1252_preserved_or_explicitly_upgraded_for_new_characters(self):
        source = "M\u00fasica: Marcos\r\n".encode("cp1252")
        data, report = self.change(source, 'Troque "Marcos" por "Mateus"')
        self.assertEqual(data, "M\u00fasica: Mateus\r\n".encode("cp1252"))
        data, report = self.change(source, 'Troque "Marcos" por "\U0001f600"')
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("\U0001f600", decode_document_text(data)[0])
        self.assertTrue(any("codificacao" in warning for warning in report["warnings"]))

    def test_multiline_replacements_keep_original_line_endings_and_surroundings(self):
        for ending in ("\n", "\r\n", "\r"):
            with self.subTest(ending=repr(ending)):
                source = ending.join(["Antes", "Marcos", "Depois", ""]).encode()
                data, _ = self.change(
                    source, 'Troque "Marcos" por "Mateus\nNova linha"'
                )
                self.assertEqual(
                    data,
                    ending.join(
                        ["Antes", "Mateus", "Nova linha", "Depois", ""]
                    ).encode(),
                )

    def test_all_first_second_last_and_invalid_occurrence(self):
        source = b"Marcos, Marcos\nMarcos\nMantido\n"
        cases = [
            ("todas as ocorrencias de", b"Ana, Ana\nAna\nMantido\n"),
            ("a primeira ocorrencia de", b"Ana, Marcos\nMarcos\nMantido\n"),
            ("a segunda ocorrencia de", b"Marcos, Ana\nMarcos\nMantido\n"),
            ("a ultima ocorrencia de", b"Marcos, Marcos\nAna\nMantido\n"),
        ]
        for scope, expected in cases:
            with self.subTest(scope=scope):
                self.assertEqual(
                    self.change(source, f'Troque {scope} "Marcos" por "Ana"')[0],
                    expected,
                )
        with self.assertRaisesRegex(DocumentEditError, "nao existe"):
            self.change(source, 'Troque a 4a ocorrencia de "Marcos" por "Ana"')

    def test_unknown_quoted_target_cannot_hallucinate_a_rewrite(self):
        with self.assertRaisesRegex(DocumentEditError, "nao foi encontrado"):
            self.change(b"Nome: Marcos", 'Troque "Nome: Roberto" por "Ana"')

    def test_randomized_all_replacements_match_literal_reference(self):
        rng = random.Random(230)
        for _ in range(150):
            text = "\n".join(
                " ".join(rng.choices(["Marcos", "Outros", "Mantidos"], k=10))
                for _ in range(3)
            )
            if "Marcos" in text:
                data, _ = self.change(
                    text.encode(), 'Troque todas as ocorrencias de "Marcos" por "Ana"'
                )
                self.assertEqual(data.decode(), text.replace("Marcos", "Ana"))

    def test_binary_and_invalid_declared_encoding_are_not_silently_replaced(self):
        for data in (b"abc\0def", b"\xff\xfea"):
            with self.assertRaises(GeneratedFileError):
                inspect_document(data, "txt")

    def test_utf16_json_can_be_edited_and_validated(self):
        source = b"\xff\xfe" + '{"nome":"Marcos"}'.encode("utf-16-le")
        snapshot = inspect_document(source, "json")
        data, _ = apply_document_edits(
            snapshot, quoted_replacement(snapshot, 'Troque "Marcos" por "Ana"')
        )
        validate_document_output(data, "json")
        self.assertEqual(json.loads(decode_document_text(data)[0]), {"nome": "Ana"})

    def test_python_utf16_is_converted_to_executable_utf8_with_warning(self):
        source = b"\xff\xfe" + 'nome = "Marcos"\r\n'.encode("utf-16-le")
        snapshot = inspect_document(source, "py")
        data, report = apply_document_edits(
            snapshot, quoted_replacement(snapshot, 'Troque "Marcos" por "Ana"')
        )
        validate_document_output(data, "py")
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertTrue(any("codificacao" in warning for warning in report["warnings"]))

    def test_xml_encoding_declaration_is_updated_when_emoji_requires_utf8(self):
        source = (
            '<?xml version="1.0" encoding="windows-1252"?><doc>Marcos</doc>'.encode(
                "cp1252"
            )
        )
        snapshot = inspect_document(source, "xml")
        data, report = apply_document_edits(
            snapshot, quoted_replacement(snapshot, 'Troque "Marcos" por "\U0001f600"')
        )
        validate_document_output(data, "xml")
        self.assertIn('encoding="UTF-8"', decode_document_text(data)[0])
        self.assertTrue(any("codificacao" in warning for warning in report["warnings"]))


class OfficeStyleScenarios(unittest.TestCase):
    def test_literal_spreadsheet_edits_keep_numeric_boolean_formula_types(self):
        from openpyxl import Workbook, load_workbook

        for objective, cell, value, kind in [
            ('Troque "10" por "20"', "A1", 20, "n"),
            ('Troque "1" por "0"', "B1", False, "b"),
            ('Troque "=SUM(A1,10)" por "=SUM(A1,20)"', "C1", "=SUM(A1,20)", "f"),
            ('Troque "=texto" por "=literal"', "D1", "=literal", "s"),
        ]:
            with self.subTest(cell=cell):
                document = Workbook()
                document.active[cell] = {
                    "A1": 10,
                    "B1": True,
                    "C1": "=SUM(A1,10)",
                    "D1": "=texto",
                }[cell]
                if cell == "D1":
                    document.active[cell].data_type = "s"
                snapshot = inspect_document(saved(document), "xlsx")
                data, _ = apply_document_edits(
                    snapshot, quoted_replacement(snapshot, objective)
                )
                workbook = load_workbook(io.BytesIO(data), data_only=False)
                self.assertEqual(workbook.active[cell].value, value)
                self.assertEqual(workbook.active[cell].data_type, kind)
                validate_document_output(data, "xlsx")

    def test_docx_partial_style_across_runs_does_not_affect_surrounding_text(self):
        from docx import Document

        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Antes Ma").italic = True
        paragraph.add_run("rcos Depois")
        snapshot = inspect_document(saved(document), "docx")
        element = snapshot.elements[0]
        data, _ = apply_document_edits(
            snapshot,
            [
                {
                    "id": element["id"],
                    "old_text": element["text"],
                    "new_text": element["text"],
                    "style": {"bold": True, "color": "112233"},
                    "style_text": "Marcos",
                }
            ],
        )
        result = Document(io.BytesIO(data)).paragraphs[0]
        self.assertEqual(result.text, "Antes Marcos Depois")
        self.assertEqual(
            [(r.text, r.bold, r.italic) for r in result.runs],
            [
                ("Antes ", None, True),
                ("Ma", True, True),
                ("rcos", True, None),
                (" Depois", None, None),
            ],
        )
        validate_document_output(data, "docx")

    def test_pptx_partial_style_and_repeated_word_occurrence(self):
        from pptx import Presentation

        presentation = Presentation()
        shape = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        ).shapes.title
        shape.text = "Marcos e Marcos"
        snapshot = inspect_document(saved(presentation), "pptx")
        item = snapshot.elements[0]
        edit = {
            "id": item["id"],
            "old_text": item["text"],
            "new_text": item["text"],
            "style": {"italic": True},
            "style_text": "Marcos",
        }
        with self.assertRaisesRegex(DocumentEditError, "ambiguo"):
            apply_document_edits(snapshot, [edit])
        data, _ = apply_document_edits(snapshot, [dict(edit, style_occurrence=2)])
        runs = (
            Presentation(io.BytesIO(data))
            .slides[0]
            .shapes.title.text_frame.paragraphs[0]
            .runs
        )
        self.assertEqual(
            [(r.text, r.font.italic) for r in runs],
            [("Marcos e ", None), ("Marcos", True)],
        )
        validate_document_output(data, "pptx")

    def test_partial_alignment_and_missing_style_are_rejected(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Marcos")
        snapshot = inspect_document(saved(document), "docx")
        element = snapshot.elements[0]
        base = {
            "id": element["id"],
            "old_text": "Marcos",
            "new_text": "Marcos",
            "style_text": "Marcos",
        }
        for extra in ({"style": {"alignment": "center"}}, {}):
            with self.assertRaises(DocumentEditError):
                apply_document_edits(snapshot, [dict(base, **extra)])


class PlannerFailureScenarios(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_json_gets_one_bounded_repair(self):
        snapshot = inspect_document(b"Nome: Marcos", "txt")
        complete = AsyncMock(
            side_effect=[
                json.JSONDecodeError("truncated", "{", 1),
                {
                    "edits": [
                        {
                            "id": "line1",
                            "old_text": "Nome: Marcos",
                            "new_text": "Nome: Ana",
                            "value_type": "text",
                        }
                    ]
                },
            ]
        )
        edits = await propose_document_edits(
            snapshot, "Corrija o nome para Ana", complete
        )
        self.assertEqual(edits[0]["new_text"], "Nome: Ana")
        self.assertEqual(complete.await_count, 2)

    async def test_wrong_json_type_and_repeated_failure_are_explicit(self):
        snapshot = inspect_document(b"Nome: Marcos", "txt")
        for answer in ([], "invalid", {"edits": "invalid"}):
            complete = AsyncMock(return_value=answer)
            with self.assertRaises(DocumentEditError):
                await propose_document_edits(snapshot, "Corrija o nome", complete)
            self.assertEqual(complete.await_count, 2)

    async def test_cancellation_is_not_retried_or_swallowed(self):
        complete = AsyncMock(side_effect=asyncio.CancelledError())
        with self.assertRaises(asyncio.CancelledError):
            await propose_document_edits(
                inspect_document(b"Nome: Marcos", "txt"), "Corrija o nome", complete
            )
        complete.assert_awaited_once()


class CreationFormatScenarios(unittest.TestCase):
    def test_duplicate_json_keys_and_nonfinite_numbers_are_rejected(self):
        for content in (
            '{"nome":"Marcos","nome":"Ana"}',
            '{"valor":NaN}',
            '{"valor":Infinity}',
        ):
            with self.subTest(content=content):
                with self.assertRaises(GeneratedFileError):
                    build_generated_file("test.json", content, "json")
                with self.assertRaises(GeneratedFileError):
                    validate_document_output(content.encode(), "json")

    def test_preferred_office_engine_and_fallback_generate_valid_documents(self):
        fixtures = {
            "docx": "# Titulo\n\nParagrafo com acentos: m\u00fasica.\n\n| Nome | Valor |\n| --- | --- |\n| Mateus | 10 |",
            "xlsx": json.dumps(
                {
                    "sheets": [
                        {"name": "Dados", "rows": [["Nome", "Valor"], ["Mateus", 10]]}
                    ]
                }
            ),
            "pptx": json.dumps(
                {
                    "slides": [
                        {
                            "title": "Titulo",
                            "content": ["Primeiro ponto", "Segundo ponto"],
                        }
                    ]
                }
            ),
        }
        for fmt, content in fixtures.items():
            with self.subTest(format=fmt):
                _, data, _ = build_generated_file(
                    "test." + fmt, content, fmt, prefer_officecli=True
                )
                validate_document_output(data, fmt, content)

    def test_pdf_emoji_fallback_preserves_character_in_output(self):
        from pathlib import Path
        from pypdf import PdfReader

        if not Path("C:/Windows/Fonts/seguiemj.ttf").is_file():
            self.skipTest("System emoji font unavailable")
        _, data, _ = build_generated_file("test.pdf", "Musica \U0001f600", "pdf")
        self.assertIn("\U0001f600", PdfReader(io.BytesIO(data)).pages[0].extract_text())
        validate_document_output(data, "pdf")

    def test_unavailable_pdf_glyph_is_explicit_not_dropped(self):
        with self.assertRaisesRegex(GeneratedFileError, "caracteres"):
            build_generated_file("test.pdf", "Texto \U0010fffd", "pdf")

    def test_every_advertised_output_format_creates_and_validates(self):
        special = {
            "json": '{"nome":"Mateus","valor":1}',
            "xml": "<doc><nome>Mateus</nome></doc>",
            "yaml": "nome: Mateus\nvalor: 1",
            "yml": "nome: Mateus\nvalor: 1",
            "py": 'def nome():\n    return "Mateus"\n',
            "xlsx": json.dumps(
                {
                    "sheets": [
                        {"name": "Dados", "rows": [["Nome", "Valor"], ["Mateus", 1]]}
                    ]
                }
            ),
            "pptx": json.dumps(
                {"slides": [{"title": "Mateus", "content": ["Resumo", "Conteudo"]}]}
            ),
            "zip": json.dumps(
                {"files": [{"path": "pasta/mateus.txt", "content": "Mateus"}]}
            ),
            "csv": "Nome,Valor\nMateus,1",
            "srt": "1\n00:00:00,000 --> 00:00:01,000\nMateus\n",
        }
        for fmt in MIME_TYPES:
            with self.subTest(format=fmt):
                content = special.get(fmt, "Mateus\nConteudo de teste")
                name, data, mime = build_generated_file(f"document.{fmt}", content, fmt)
                self.assertTrue(name.endswith("." + fmt))
                self.assertEqual(mime, MIME_TYPES[fmt])
                self.assertTrue(data)
                self.assertEqual(
                    validate_document_output(data, fmt, content)["validation"], "passed"
                )

    def test_invalid_python_is_not_published_as_valid_code(self):
        with self.assertRaisesRegex(GeneratedFileError, "sintaxe"):
            validate_document_output(b"def invalid(\n", "py")

    def test_utf16_entities_cannot_bypass_xml_security_check(self):
        value = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE doc [<!ENTITY x "unsafe">]><doc>&x;</doc>'
        with self.assertRaises(DocumentEditError):
            validate_document_output(b"\xff\xfe" + value.encode("utf-16-le"), "xml")


if __name__ == "__main__":
    unittest.main()
